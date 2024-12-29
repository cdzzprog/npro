# def remove_empty_lines(input_file, output_file):
#     try:
#         # 打开并读取输入文件
#         with open(input_file, 'r', encoding='utf-8') as infile:
#             lines = infile.readlines()
        
#         # 过滤掉空行，空行是指去除前后空格后的内容为空的行
#         non_empty_lines = [line for line in lines if line.strip() != '']
        
#         # 打开输出文件，写入去除空行后的内容
#         with open(output_file, 'w', encoding='utf-8') as outfile:
#             outfile.writelines(non_empty_lines)
        
#         print(f"处理完成，空行已移除，结果保存为 '{output_file}'")
    
#     except Exception as e:
#         print(f"处理文件时发生错误: {e}")








# # 调用函数，输入和输出文件的路径
# input_file =r'C:\Users\龙儿璨\Desktop\d.txt'  # 输入文件路径（原始文件）
# output_file = r'C:\Users\龙儿璨\Desktop\output.txt'  # 输出文件路径（处理后文件）

# # 执行删除空行操作
# remove_empty_lines(input_file, output_file)

from docx import Document

def check_empty_lines(doc_path):
    # 打开文档
    doc = Document(doc_path)
    
    empty_lines = []  # 存储空行的索引
    total_lines = len(doc.paragraphs)  # 获取文档中的总行数（段落数）
    
    for i, para in enumerate(doc.paragraphs):
        # 检查段落是否为空或仅包含空格
        if not para.text.strip():  # strip() 会去掉前后空白字符
            empty_lines.append(i)
    
    empty_line_count = len(empty_lines)  # 计算空行的数量
    non_empty_line_count = total_lines - empty_line_count  # 计算非空行的数量
    
    print(f"文档总共有 {total_lines} 行（段落）。")
    print(f"其中有 {empty_line_count} 行为空行。")
    print(f"其中有 {non_empty_line_count} 行是非空行。")
    
    if empty_lines:
        print(f"空行的索引为：{empty_lines}")
    else:
        print("文档中没有空行")

# 示例用法
doc_path = r'E:\papers\zip\软著模板\高原地区山体滑坡遥感智能识别系统V1.0-源码.docx'  # 替换为实际的文档路径
check_empty_lines(doc_path)
